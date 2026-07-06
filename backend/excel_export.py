# backend/excel_export.py
# One-click session workbook: every metric, table, and a rendered path map
# in a single styled .xlsx file.  Researchers do cross-session comparison
# from these workbooks, so completeness and consistent layout matter.
#
# Sheets:
#   Summary          — video info, run settings, headline metrics
#   Path Map         — rendered path (and heatmap) over blueprint / white
#   Episodes         — active/stationary episode table
#   Regions          — per-region dwell / visits / latency
#   Region Segments  — every entry→exit presence segment
#   Speed            — full speed-vs-time series
#   Path Points      — the smoothed path samples used for all metrics
#
# The path image is drawn with matplotlib (Agg) and embedded via openpyxl.
# The blueprint image, when provided, is assumed to be a top-down view of
# the room oriented like the in-app map: camera side at the BOTTOM.
#
# No PyQt6 imports — pure backend, callable from anywhere.

from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path

import numpy as np

# ---------------------------------------------------------------------------
# Styling constants
# ---------------------------------------------------------------------------

_HEADER_FILL   = "1F3864"   # dark blue
_HEADER_FONT   = "FFFFFF"
_TITLE_FILL    = "2E5395"
_ALT_ROW_FILL  = "EDF2F9"   # light blue-grey for zebra striping
_ACTIVE_FONT   = "1E7B1E"   # green
_STATION_FONT  = "9C3333"   # red


def _styles():
    """Build the reusable openpyxl style objects (import-safe helper)."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    thin   = Side(style="thin", color="B0B8C4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    return {
        "title_font":  Font(bold=True, size=15, color=_HEADER_FONT),
        "title_fill":  PatternFill("solid", fgColor=_TITLE_FILL),
        "header_font": Font(bold=True, size=11, color=_HEADER_FONT),
        "header_fill": PatternFill("solid", fgColor=_HEADER_FILL),
        "section_font": Font(bold=True, size=12, color=_HEADER_FILL),
        "alt_fill":    PatternFill("solid", fgColor=_ALT_ROW_FILL),
        "border":      border,
        "center":      Alignment(horizontal="center", vertical="center"),
        "left":        Alignment(horizontal="left",  vertical="center"),
    }


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def export_session_excel(
    xlsx_path: str | Path,
    video_name: str,
    stats,                        # PathStats
    ep_sum: dict,                 # episode_summary() output
    coverage_pct: float,
    episodes: list,               # list[MovementEpisode]
    reg_stats: list,              # list[RegionStats]
    region_segments: list[dict],  # {"start_ms","end_ms","region","frames"}
    speed_series: list,           # list[SpeedSample]
    path_points: list,            # SMOOTHED list[PathPoint]
    room_size_cm: tuple,
    analysis_params: dict,        # smooth_window / stationary threshold / min episode
    run_info: dict | None = None, # yolo model, stride, dates …
    blueprint_path: str | None = None,
    region_shapes: list[dict] | None = None,  # from project_region_shapes()
) -> None:
    """Write the full session workbook to *xlsx_path*."""
    from openpyxl import Workbook

    wb = Workbook()
    st = _styles()

    _write_summary_sheet(
        wb.active, st, video_name, stats, ep_sum, coverage_pct,
        room_size_cm, analysis_params, run_info or {},
        n_regions=len(reg_stats),
    )
    _write_path_map_sheet(
        wb.create_sheet("Path Map"), st,
        path_points, room_size_cm, stats.calibrated, blueprint_path,
        region_shapes or [],
    )
    _write_episodes_sheet(wb.create_sheet("Episodes"), st, episodes)
    _write_regions_sheet(wb.create_sheet("Regions"), st, reg_stats)
    _write_region_segments_sheet(
        wb.create_sheet("Region Segments"), st, region_segments)
    _write_speed_sheet(wb.create_sheet("Speed"), st, speed_series)
    _write_path_points_sheet(wb.create_sheet("Path Points"), st, path_points)

    wb.save(str(xlsx_path))


# ---------------------------------------------------------------------------
# Sheet writers
# ---------------------------------------------------------------------------

def _write_table(ws, st, header: list[str], rows: list[list], start_row: int = 1):
    """Write a zebra-striped, bordered table with a styled header row."""
    for col, text in enumerate(header, start=1):
        c = ws.cell(row=start_row, column=col, value=text)
        c.font, c.fill  = st["header_font"], st["header_fill"]
        c.border, c.alignment = st["border"], st["center"]
    for r, row in enumerate(rows, start=start_row + 1):
        for col, val in enumerate(row, start=1):
            c = ws.cell(row=r, column=col, value=val)
            c.border = st["border"]
            if (r - start_row) % 2 == 0:
                c.fill = st["alt_fill"]
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    _autofit(ws)


def _autofit(ws, min_w: int = 10, max_w: int = 42):
    from openpyxl.utils import get_column_letter
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        width  = max(
            (len(str(c.value)) for c in col_cells if c.value is not None),
            default=0,
        )
        ws.column_dimensions[letter].width = max(min_w, min(max_w, width + 3))


def _write_summary_sheet(ws, st, video_name, stats, ep_sum, coverage_pct,
                         room_size_cm, analysis_params, run_info, n_regions):
    ws.title = "Summary"

    # Title banner
    ws.merge_cells("A1:D1")
    c = ws["A1"]
    c.value = f"Playroom Session Report — {video_name}"
    c.font, c.fill, c.alignment = st["title_font"], st["title_fill"], st["center"]
    ws.row_dimensions[1].height = 28

    ws["A2"] = f"Exported {datetime.now().strftime('%Y-%m-%d %H:%M')}"

    def _section(row: int, label: str) -> int:
        cell = ws.cell(row=row, column=1, value=label)
        cell.font = st["section_font"]
        return row + 1

    def _kv(row: int, pairs: list[tuple]) -> int:
        for key, val in pairs:
            k = ws.cell(row=row, column=1, value=key)
            v = ws.cell(row=row, column=2, value=val)
            k.border = v.border = st["border"]
            k.font = _bold()
            row += 1
        return row + 1

    cal = stats.calibrated
    row = 4
    row = _section(row, "Session")
    row = _kv(row, [
        ("Video",            video_name),
        ("Duration",         _fmt_dur_s(stats.duration_s)),
        ("Track points",     stats.n_points),
        ("Interpolated",     stats.n_interpolated),
        ("Calibrated",       "yes" if cal else "NO — distances/speeds unavailable"),
        ("Room size",        f"{room_size_cm[0]/100:.2f} m × {room_size_cm[1]/100:.2f} m"),
        ("Regions defined",  n_regions),
    ])

    row = _section(row, "Movement Metrics")
    if cal:
        row = _kv(row, [
            ("Total distance (m)",       round(stats.total_distance_m, 2)),
            ("Average speed (m/s)",      round(stats.avg_speed_m_s, 3)),
            ("Floor coverage (%)",       round(coverage_pct, 1)),
            ("Active time (%)",          round(ep_sum.get("active_ratio", 0) * 100, 1)),
            ("Stationary time (%)",      round(100 - ep_sum.get("active_ratio", 0) * 100, 1)),
            ("Active episodes",          ep_sum.get("n_active_episodes", 0)),
            ("Stationary episodes",      ep_sum.get("n_stationary_episodes", 0)),
            ("Avg active episode (s)",   round(ep_sum.get("avg_active_dur_s", 0), 1)),
            ("Avg stationary episode (s)", round(ep_sum.get("avg_stationary_dur_s", 0), 1)),
        ])
    else:
        row = _kv(row, [("Movement metrics", "(no calibration)")])

    row = _section(row, "Analysis Parameters")
    row = _kv(row, [(k, v) for k, v in analysis_params.items()])

    if run_info:
        row = _section(row, "Processing")
        row = _kv(row, [(k, str(v)) for k, v in run_info.items()])

    _autofit(ws, min_w=24)


def _write_path_map_sheet(ws, st, path_points, room_size_cm,
                          calibrated, blueprint_path, region_shapes):
    from openpyxl.drawing.image import Image as XLImage

    ws["A1"] = "Path Map"
    ws["A1"].font = st["section_font"]
    note = ("Path over room blueprint" if blueprint_path
            else "Path (no blueprint configured — plain background)")
    if not calibrated:
        note += "  ·  UNCALIBRATED: pixel space, not to scale"
    ws["A2"] = note

    # Plain path (no regions)
    png = _render_path_png(path_points, room_size_cm, calibrated,
                           blueprint_path, with_heatmap=False)
    if png is not None:
        ws.add_image(XLImage(png), "A4")

    # Path with region footprints overlaid
    if region_shapes:
        pr = _render_path_png(path_points, room_size_cm, calibrated,
                              blueprint_path, with_heatmap=False,
                              region_shapes=region_shapes)
        if pr is not None:
            ws["L1"] = "Path + Regions"
            ws["L1"].font = st["section_font"]
            ws.add_image(XLImage(pr), "L4")

    if calibrated:   # heatmap needs real-world coordinates
        hm = _render_path_png(path_points, room_size_cm, calibrated,
                              blueprint_path, with_heatmap=True)
        if hm is not None:
            anchor_col = "W" if region_shapes else "L"
            ws[f"{anchor_col}1"] = "Heatmap"
            ws[f"{anchor_col}1"].font = st["section_font"]
            ws.add_image(XLImage(hm), f"{anchor_col}4")


def _write_episodes_sheet(ws, st, episodes):
    header = ["#", "Start", "End", "Duration (s)", "Type",
              "Distance (m)", "Avg Speed (m/s)"]
    rows = []
    for i, ep in enumerate(episodes, start=1):
        rows.append([
            i, _fmt_ms(ep.start_ms), _fmt_ms(ep.end_ms),
            round(ep.duration_s, 2), ep.kind,
            round(ep.distance_m, 3) if ep.kind == "active" else None,
            round(ep.avg_speed_m_s, 3) if ep.kind == "active" else None,
        ])
    _write_table(ws, st, header, rows)
    # Colour the Type column
    from openpyxl.styles import Font
    for r, ep in enumerate(episodes, start=2):
        color = _ACTIVE_FONT if ep.kind == "active" else _STATION_FONT
        ws.cell(row=r, column=5).font = Font(bold=True, color=color)


def _write_regions_sheet(ws, st, reg_stats):
    header = ["Region", "Dwell (s)", "Dwell (%)", "Visits",
              "First Visit (s)"]
    rows = []
    for rs in reg_stats:
        rows.append([
            rs.region_name,
            round(rs.total_dwell_s, 1),
            round(rs.dwell_pct, 1),
            rs.visit_count,
            round(rs.first_visit_ms / 1000.0, 1)
            if rs.first_visit_ms is not None else None,
        ])
    _write_table(ws, st, header, rows)


def _write_region_segments_sheet(ws, st, segments):
    header = ["#", "Region", "Start", "End", "Duration (s)", "Samples"]
    rows = []
    for i, s in enumerate(segments, start=1):
        rows.append([
            i, s["region"], _fmt_ms(s["start_ms"]), _fmt_ms(s["end_ms"]),
            round((s["end_ms"] - s["start_ms"]) / 1000.0, 2),
            s.get("frames"),
        ])
    _write_table(ws, st, header, rows)


def _write_speed_sheet(ws, st, speed_series):
    header = ["Time (s)", "Speed (m/s)", "Speed (cm/s)"]
    rows = [
        [round(s.timestamp_ms / 1000.0, 2),
         round(s.speed_m_s, 4), round(s.speed_cm_s, 2)]
        for s in speed_series
    ]
    _write_table(ws, st, header, rows)


def _write_path_points_sheet(ws, st, path_points):
    header = ["Frame", "Time (s)", "Pixel X", "Pixel Y",
              "Floor X (cm)", "Floor Y (cm)", "Interpolated"]
    rows = [
        [p.frame_index, round(p.timestamp_ms / 1000.0, 3),
         round(p.cx_px, 1), round(p.cy_px, 1),
         round(p.x_cm, 1) if p.x_cm is not None else None,
         round(p.y_cm, 1) if p.y_cm is not None else None,
         int(p.interpolated)]
        for p in path_points
    ]
    _write_table(ws, st, header, rows)


# ---------------------------------------------------------------------------
# Path map rendering (matplotlib Agg → PNG BytesIO)
# ---------------------------------------------------------------------------

# Region palette — matches the UI overlay colours
_REGION_COLORS = [
    "#FF6464", "#64DC64", "#648CFF", "#FFDC32",
    "#FF64DC", "#50DCDC", "#FFA028", "#B464FF",
]


def _render_path_png(
    path_points: list,
    room_size_cm: tuple,
    calibrated: bool,
    blueprint_path: str | None,
    with_heatmap: bool,
    region_shapes: list[dict] | None = None,
) -> io.BytesIO | None:
    """
    Render the path (blue=early → red=late) on the blueprint or a white
    background.  World space when calibrated (camera at bottom), pixel space
    otherwise.  region_shapes (from project_region_shapes) draws each
    region's floor footprint.  Returns a PNG BytesIO or None when empty.
    """
    if not path_points:
        return None

    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.collections import LineCollection

    if calibrated:
        pts = [(p.x_cm, p.y_cm) for p in path_points
               if p.x_cm is not None and p.y_cm is not None]
        if not pts:
            return None
        xs, ys = zip(*pts)
        room_w, room_d = room_size_cm
        extent = (0, room_w, 0, room_d)      # y up = away from camera
        aspect_w, aspect_h = room_w, room_d
    else:
        xs = [p.cx_px for p in path_points]
        ys = [p.cy_px for p in path_points]
        # Pixel y grows downward — flip so the plot isn't mirrored
        ymax = max(ys)
        ys   = [ymax - y for y in ys]
        extent = (min(xs), max(xs) or 1, 0, max(ys) or 1)
        aspect_w = (extent[1] - extent[0]) or 1
        aspect_h = (extent[3] - extent[2]) or 1

    fig_h = 7.0
    fig_w = max(3.0, min(10.0, fig_h * aspect_w / aspect_h))
    fig, ax = plt.subplots(figsize=(fig_w, fig_h), dpi=110)

    # Background: blueprint or white
    if blueprint_path and Path(blueprint_path).exists():
        try:
            from PIL import Image as PILImage
            bp = PILImage.open(blueprint_path)
            ax.imshow(bp, extent=extent, origin="upper",
                      aspect="auto", alpha=0.9, zorder=0)
        except Exception:
            ax.set_facecolor("white")
    else:
        ax.set_facecolor("white")

    # Region floor footprints (drawn under the path)
    if region_shapes and calibrated:
        from matplotlib.patches import Polygon as MPolygon
        for idx, shape in enumerate(region_shapes):
            color = _REGION_COLORS[idx % len(_REGION_COLORS)]
            if shape["polygon_cm"]:
                poly = MPolygon(
                    shape["polygon_cm"], closed=True,
                    facecolor=color, edgecolor=color,
                    alpha=0.35, linewidth=1.6, linestyle="--", zorder=1.5,
                )
                ax.add_patch(poly)
                cx = sum(p[0] for p in shape["polygon_cm"]) / 4
                cy = sum(p[1] for p in shape["polygon_cm"]) / 4
            else:
                cx, cy = shape["point_cm"]
                ax.scatter([cx], [cy], c=color, s=90, zorder=1.5,
                           edgecolors="black")
            room_w, room_d = room_size_cm
            cx = min(max(cx, 2), room_w - 2)
            cy = min(max(cy, 2), room_d - 2)
            ax.text(cx, cy, shape["name"], fontsize=8, weight="bold",
                    ha="center", va="center", color="black", zorder=4,
                    bbox=dict(boxstyle="round,pad=0.25",
                              facecolor="white", alpha=0.75, linewidth=0))

    # Heatmap overlay
    if with_heatmap and calibrated:
        from backend.path_analyzer import compute_heatmap
        grid = compute_heatmap(path_points, room_size_cm)
        if grid.max() > 0:
            masked = np.ma.masked_where(grid < 0.02, grid)
            ax.imshow(masked, extent=extent, origin="upper",
                      cmap="jet", alpha=0.55, aspect="auto", zorder=1)

    # Path polyline, blue → red along time
    if not with_heatmap or not calibrated:
        segs = [((xs[i - 1], ys[i - 1]), (xs[i], ys[i]))
                for i in range(1, len(xs))]
        if segs:
            lc = LineCollection(
                segs, cmap="coolwarm",
                array=np.linspace(0, 1, len(segs)),
                linewidths=1.6, zorder=2,
            )
            ax.add_collection(lc)
        ax.scatter([xs[0]], [ys[0]], c="#22bb55", s=70, zorder=3,
                   edgecolors="black", label="Start")
        ax.scatter([xs[-1]], [ys[-1]], c="#dd3333", s=70, zorder=3,
                   edgecolors="black", label="End")
        ax.legend(loc="upper right", fontsize=8)

    ax.set_xlim(extent[0], extent[1])
    ax.set_ylim(extent[2], extent[3])
    if calibrated:
        ax.set_xlabel("Room width (cm) — camera at bottom", fontsize=8)
        ax.set_ylabel("Room depth (cm)", fontsize=8)
    else:
        ax.set_xticks([]); ax.set_yticks([])
    ax.tick_params(labelsize=7)
    ax.set_title("Heatmap" if with_heatmap else "Movement Path", fontsize=10)

    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight",
                facecolor="white")
    plt.close(fig)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

def _bold():
    from openpyxl.styles import Font
    return Font(bold=True)


def _fmt_ms(ms: float) -> str:
    from backend.path_analyzer import _fmt_ms as fmt
    return fmt(ms)


def _fmt_dur_s(seconds: float) -> str:
    m = int(seconds) // 60
    s = seconds % 60
    return f"{m}m {s:.0f}s" if m else f"{s:.1f}s"
